from pathlib import Path

from lxml import etree
from pptx import Presentation


root = Path(__file__).resolve().parents[1]
figures = root / "manuscript_temporal" / "figures"
svg_path = figures / "fig_system_workflow.svg"
pptx_path = figures / "fig_system_workflow_editable.pptx"

# SVG text index, PowerPoint shape index, replacement lines. The two source
# files contain the same diagram as editable text and shapes.
edits = [
    (29, 70, ["Clean: 1, 2, 4, 5", "Pairs: 1–2, 4–5"]),
    (33, 80, ["Adjacent pairs", "Δindex = 1"]),
    (39, 92, ["Adjacent pairs", "Δindex = 1"]),
    (50, 136, ["Train scores"]),
    (51, 140, ["Train median"]),
    (52, 141, ["Test subjects"]),
    (54, 145, ["Source scores"]),
    (55, 149, ["Source median"]),
    (56, 150, ["External subjects"]),
]

tree = etree.parse(str(svg_path))
tree.getroot().set('width', '183mm')
tree.getroot().set('height', '102.995mm')
texts = tree.xpath('//s:text', namespaces={'s': 'http://www.w3.org/2000/svg'})
presentation = Presentation(str(pptx_path))
shapes = presentation.slides[0].shapes
for text_index, shape_index, lines in edits:
    spans = texts[text_index].xpath('./s:tspan', namespaces={'s': 'http://www.w3.org/2000/svg'})
    paragraphs = shapes[shape_index].text_frame.paragraphs
    assert len(spans) == len(paragraphs) == len(lines)
    for span, paragraph, line in zip(spans, paragraphs, lines):
        assert len(paragraph.runs) == 1
        span.text = line
        paragraph.runs[0].text = line

tree.write(str(svg_path), encoding='utf-8', xml_declaration=True)
presentation.save(str(pptx_path))
print('Updated SVG and editable PowerPoint labels.')
